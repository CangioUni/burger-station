import os
import datetime
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from fpdf import FPDF

class StatsExporter:
    def __init__(self, session_factory):
        self.session_factory = session_factory

    def _get_models(self):
        from main import Order, OrderItem, MenuItem, Category
        return Order, OrderItem, MenuItem, Category

    def get_available_days(self) -> list[str]:
        """Returns a list of unique days (YYYY-MM-DD) that have orders, sorted descending."""
        Order, _, _, _ = self._get_models()
        db = self.session_factory()
        days_set = set()
        try:
            orders = db.query(Order.timestamp).all()
            for (ts,) in orders:
                if not ts:
                    continue
                if isinstance(ts, (datetime.datetime, datetime.date)):
                    days_set.add(ts.strftime("%Y-%m-%d"))
                elif isinstance(ts, str):
                    day_str = ts.strip()[:10]
                    if len(day_str) == 10 and day_str[4] == '-' and day_str[7] == '-':
                        days_set.add(day_str)
        finally:
            db.close()

        # Sort descending (most recent first)
        return sorted(list(days_set), reverse=True)

    def generate_stats(self, day: str) -> dict:
        """Computes comprehensive daily analytics for a given day (YYYY-MM-DD)."""
        Order, OrderItem, MenuItem, Category = self._get_models()
        db = self.session_factory()
        
        try:
            # Query all menu items to build category & burger lookups
            all_menu_items = db.query(MenuItem).all()
            item_cat_map = {}
            combo_set = set()
            burger_set = set()

            for mi in all_menu_items:
                desc_clean = mi.description.strip().lower()
                cat_clean = (mi.category or "").strip().lower()
                item_cat_map[desc_clean] = cat_clean
                if mi.is_combo or cat_clean in ['menu', 'menù', 'combo']:
                    combo_set.add(desc_clean)
                if cat_clean in ['panini', 'panino', 'burger', 'burgers', 'hamburger'] or 'panino' in desc_clean or 'burger' in desc_clean:
                    burger_set.add(desc_clean)

            # Query all orders
            all_orders = db.query(Order).all()
            day_orders = []

            for o in all_orders:
                ts = o.timestamp
                o_day = ""
                if isinstance(ts, (datetime.datetime, datetime.date)):
                    o_day = ts.strftime("%Y-%m-%d")
                elif isinstance(ts, str):
                    o_day = ts.strip()[:10]
                
                if o_day == day:
                    day_orders.append(o)

            total_orders = len(day_orders)
            total_revenue = 0.0
            total_discount = 0.0

            revenue_by_payment = defaultdict(float)
            orders_over_time = defaultdict(int)

            total_menus = 0
            total_burgers = 0
            burger_types = defaultdict(int)
            items_by_category = defaultdict(lambda: defaultdict(int))
            orders_list = []

            for o in day_orders:
                o_total = float(o.total or 0.0)
                o_disc = float(o.discount or 0.0)
                total_revenue += o_total
                total_discount += o_disc

                pay_method = (o.payment_method or "Contanti").strip()
                if not pay_method or pay_method.lower() in ["cash", "contante", "contanti"]:
                    pay_method = "Contanti"
                elif pay_method.lower() in ["bancomat", "pos", "carta"]:
                    pay_method = "Bancomat"
                elif pay_method.lower() in ["satispay"]:
                    pay_method = "Satispay"
                revenue_by_payment[pay_method] += o_total

                # Time slot breakdown (15 min)
                ts = o.timestamp
                time_str = "00:00"
                if isinstance(ts, datetime.datetime):
                    minute_rounded = (ts.minute // 15) * 15
                    time_str = f"{ts.hour:02d}:{minute_rounded:02d}"
                elif isinstance(ts, str) and len(ts) >= 16 and (' ' in ts or 'T' in ts):
                    sep = ' ' if ' ' in ts else 'T'
                    time_part = ts.split(sep)[1][:5]
                    try:
                        hh, mm = map(int, time_part.split(':'))
                        minute_rounded = (mm // 15) * 15
                        time_str = f"{hh:02d}:{minute_rounded:02d}"
                    except Exception:
                        time_str = time_part
                orders_over_time[time_str] += 1

                order_id_display = o.order_number or str(o.id)
                time_display = ""
                if isinstance(ts, datetime.datetime):
                    time_display = ts.strftime("%H:%M")
                elif isinstance(ts, str) and len(ts) >= 16 and (' ' in ts or 'T' in ts):
                    sep = ' ' if ' ' in ts else 'T'
                    time_display = ts.split(sep)[1][:5]

                orders_list.append({
                    "id": order_id_display,
                    "time": time_display,
                    "table": o.table_number or "Nessuno",
                    "takeaway": bool(o.takeaway),
                    "payment_method": pay_method,
                    "total": o_total,
                    "discount": o_disc
                })

                # Process order items
                for item in o.items:
                    desc = item.description.strip()
                    desc_lower = desc.lower()
                    qty = item.quantity or 1
                    cat = item_cat_map.get(desc_lower, "altro")

                    # If item description looks like a combo or is marked as combo
                    is_combo_item = desc_lower in combo_set or cat in ['menu', 'menù', 'combo']
                    if is_combo_item:
                        total_menus += qty
                        items_by_category["menu"][desc] += qty
                    else:
                        items_by_category[cat][desc] += qty

                    # Check if the item itself is a burger
                    if desc_lower in burger_set or 'panino' in desc_lower or 'burger' in desc_lower:
                        total_burgers += qty
                        burger_types[desc] += qty

                    # Inspect combo choices for selected sub-items (e.g. chosen burger and chosen drink)
                    if item.combo_choices:
                        for choice in item.combo_choices.split(','):
                            choice_clean = choice.strip()
                            if not choice_clean or choice_clean == 'Nessuna Scelta':
                                continue
                            base_choice = choice_clean.split('[')[0].strip()
                            choice_lower = base_choice.lower()
                            choice_cat = item_cat_map.get(choice_lower, "altro")

                            if choice_lower in burger_set or 'panino' in choice_lower or 'burger' in choice_lower:
                                total_burgers += qty
                                burger_types[base_choice] += qty
                            
                            if choice_cat and choice_cat != 'menu':
                                items_by_category[choice_cat][base_choice] += qty

            # Format sorted orders_over_time
            sorted_time_slots = dict(sorted(orders_over_time.items()))

            # Convert items_by_category to regular dict
            formatted_cat_dict = {}
            for cat_name, items_dict in sorted(items_by_category.items()):
                formatted_cat_dict[cat_name] = dict(sorted(items_dict.items(), key=lambda x: x[1], reverse=True))

            return {
                "day": day,
                "total_orders": total_orders,
                "total_revenue": round(total_revenue, 2),
                "total_discount": round(total_discount, 2),
                "orders_over_time": sorted_time_slots,
                "revenue_by_payment": {k: round(v, 2) for k, v in revenue_by_payment.items()},
                "burger_stats": {
                    "total_menus": total_menus,
                    "total_burgers": total_burgers,
                    "burger_types": dict(sorted(burger_types.items(), key=lambda x: x[1], reverse=True))
                },
                "items_by_category": formatted_cat_dict,
                "orders_list": orders_list
            }
        finally:
            db.close()

    def export_to_excel(self, stats: dict, path: str):
        """Generates a professional, multi-tab Excel spreadsheet for daily sales analysis."""
        wb = Workbook()
        
        # Style helpers
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        sub_font = Font(name="Arial", size=11, bold=True, color="374151")
        bold_font = Font(name="Arial", size=10, bold=True)
        regular_font = Font(name="Arial", size=10)
        
        primary_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid") # Dark Blue
        accent_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")  # Blue
        gray_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
        
        thin_border = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        # ----------------------------------------------------
        # Tab 1: Riepilogo Generale
        # ----------------------------------------------------
        ws1 = wb.active
        ws1.title = "Riepilogo"
        ws1.views.sheetView[0].showGridLines = True

        ws1["A1"] = f"BURGER STATION - REPORT GIORNALIERO ({stats.get('day', '')})"
        ws1["A1"].font = title_font
        ws1.merge_cells("A1:D1")

        ws1["A3"] = "INDICATORE"
        ws1["B3"] = "VALORE"
        ws1["A3"].fill = primary_fill
        ws1["A3"].font = header_font
        ws1["B3"].fill = primary_fill
        ws1["B3"].font = header_font

        kpi_rows = [
            ("Totale Ordini Effettuati", stats.get("total_orders", 0)),
            ("Incasso Totale Complessivo", f"€ {stats.get('total_revenue', 0.0):.2f}"),
            ("Totale Sconti Concessi", f"€ {stats.get('total_discount', 0.0):.2f}"),
            ("Totale Menù Venduti", stats.get("burger_stats", {}).get("total_menus", 0)),
            ("Totale Burger Preparati", stats.get("burger_stats", {}).get("total_burgers", 0)),
        ]

        curr_r = 4
        for label, val in kpi_rows:
            ws1[f"A{curr_r}"] = label
            ws1[f"B{curr_r}"] = val
            ws1[f"A{curr_r}"].font = regular_font
            ws1[f"B{curr_r}"].font = bold_font
            ws1[f"A{curr_r}"].border = thin_border
            ws1[f"B{curr_r}"].border = thin_border
            if curr_r % 2 == 0:
                ws1[f"A{curr_r}"].fill = gray_fill
                ws1[f"B{curr_r}"].fill = gray_fill
            curr_r += 1

        curr_r += 2
        ws1[f"A{curr_r}"] = "METODO DI PAGAMENTO"
        ws1[f"B{curr_r}"] = "INCASSO (€)"
        ws1[f"C{curr_r}"] = "% SU TOTALE"
        for col in ["A", "B", "C"]:
            ws1[f"{col}{curr_r}"].fill = accent_fill
            ws1[f"{col}{curr_r}"].font = header_font
            ws1[f"{col}{curr_r}"].border = thin_border

        curr_r += 1
        tot_rev = stats.get("total_revenue", 0.0)
        for method, amt in stats.get("revenue_by_payment", {}).items():
            pct = (amt / tot_rev * 100) if tot_rev > 0 else 0.0
            ws1[f"A{curr_r}"] = method
            ws1[f"B{curr_r}"] = f"€ {amt:.2f}"
            ws1[f"C{curr_r}"] = f"{pct:.1f}%"
            ws1[f"A{curr_r}"].font = regular_font
            ws1[f"B{curr_r}"].font = bold_font
            ws1[f"C{curr_r}"].font = regular_font
            ws1[f"A{curr_r}"].border = thin_border
            ws1[f"B{curr_r}"].border = thin_border
            ws1[f"C{curr_r}"].border = thin_border
            curr_r += 1

        curr_r += 2
        ws1[f"A{curr_r}"] = "TIPOLOGIA BURGER"
        ws1[f"B{curr_r}"] = "QUANTITÀ VENDUTA"
        for col in ["A", "B"]:
            ws1[f"{col}{curr_r}"].fill = accent_fill
            ws1[f"{col}{curr_r}"].font = header_font
            ws1[f"{col}{curr_r}"].border = thin_border

        curr_r += 1
        for b_name, b_qty in stats.get("burger_stats", {}).get("burger_types", {}).items():
            ws1[f"A{curr_r}"] = b_name
            ws1[f"B{curr_r}"] = b_qty
            ws1[f"A{curr_r}"].font = regular_font
            ws1[f"B{curr_r}"].font = bold_font
            ws1[f"A{curr_r}"].border = thin_border
            ws1[f"B{curr_r}"].border = thin_border
            curr_r += 1

        # ----------------------------------------------------
        # Tab 2: Articoli per Categoria
        # ----------------------------------------------------
        ws2 = wb.create_sheet(title="Vendite Categorie")
        ws2.views.sheetView[0].showGridLines = True
        ws2["A1"] = "CATEGORIA"
        ws2["B1"] = "ARTICOLO"
        ws2["C1"] = "QUANTITÀ VENDUTA"
        for col in ["A", "B", "C"]:
            ws2[f"{col}1"].fill = primary_fill
            ws2[f"{col}1"].font = header_font
            ws2[f"{col}1"].border = thin_border

        r2 = 2
        for cat_name, items in stats.get("items_by_category", {}).items():
            for item_name, qty in items.items():
                ws2[f"A{r2}"] = cat_name.capitalize()
                ws2[f"B{r2}"] = item_name
                ws2[f"C{r2}"] = qty
                ws2[f"A{r2}"].font = regular_font
                ws2[f"B{r2}"].font = regular_font
                ws2[f"C{r2}"].font = bold_font
                ws2[f"A{r2}"].border = thin_border
                ws2[f"B{r2}"].border = thin_border
                ws2[f"C{r2}"].border = thin_border
                r2 += 1

        # ----------------------------------------------------
        # Tab 3: Lista Dettagliata Ordini
        # ----------------------------------------------------
        ws3 = wb.create_sheet(title="Dettaglio Ordini")
        ws3.views.sheetView[0].showGridLines = True
        headers_ord = ["N. ORDINE", "ORARIO", "TAVOLO", "ASPORTO", "PAGAMENTO", "SCONTO (€)", "TOTALE (€)"]
        for idx, h in enumerate(headers_ord, 1):
            col_letter = get_column_letter(idx)
            ws3[f"{col_letter}1"] = h
            ws3[f"{col_letter}1"].fill = primary_fill
            ws3[f"{col_letter}1"].font = header_font
            ws3[f"{col_letter}1"].border = thin_border

        r3 = 2
        for ord_info in stats.get("orders_list", []):
            ws3[f"A{r3}"] = f"#{ord_info['id']}"
            ws3[f"B{r3}"] = ord_info["time"]
            ws3[f"C{r3}"] = ord_info["table"]
            ws3[f"D{r3}"] = "SÌ" if ord_info["takeaway"] else "NO"
            ws3[f"E{r3}"] = ord_info["payment_method"]
            ws3[f"F{r3}"] = f"€ {ord_info['discount']:.2f}"
            ws3[f"G{r3}"] = f"€ {ord_info['total']:.2f}"

            for col_idx in range(1, 8):
                cell = ws3[f"{get_column_letter(col_idx)}{r3}"]
                cell.font = regular_font
                cell.border = thin_border
            ws3[f"G{r3}"].font = bold_font
            r3 += 1

        # Auto-adjust column widths on all sheets
        for sheet in [ws1, ws2, ws3]:
            for col in sheet.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        os.makedirs(os.path.dirname(path), exist_ok=True)
        wb.save(path)

    def export_to_pdf(self, stats: dict, path: str):
        """Generates a clean, modern PDF summary report for daily operations."""
        pdf = FPDF(orientation='P', unit='mm', format='A4')
        pdf.set_auto_page_break(auto=True, margin=15)
        pdf.add_page()
        
        # Header title
        pdf.set_font("Helvetica", "B", 18)
        pdf.set_text_color(30, 58, 138) # Dark blue
        pdf.cell(0, 10, "BURGER STATION - REPORT GIORNALIERO", ln=True, align="C")
        
        pdf.set_font("Helvetica", "", 12)
        pdf.set_text_color(107, 114, 128) # Gray
        pdf.cell(0, 6, f"Data di riferimento: {stats.get('day', '')}", ln=True, align="C")
        pdf.ln(6)

        # 4 KPI boxes (2x2 grid)
        box_w = 88
        box_h = 22
        
        kpis = [
            ("TOTALE ORDINI", str(stats.get("total_orders", 0))),
            ("INCASSO TOTALE", f"EUR {stats.get('total_revenue', 0.0):.2f}"),
            ("TOTALE BURGER", str(stats.get("burger_stats", {}).get("total_burgers", 0))),
            ("MENU VENDUTI", str(stats.get("burger_stats", {}).get("total_menus", 0)))
        ]

        # First row of KPI boxes
        y_pos = pdf.get_y()
        pdf.set_fill_color(239, 246, 255) # Light blue
        pdf.set_draw_color(191, 219, 254)
        pdf.rect(14, y_pos, box_w, box_h, style='FD')
        pdf.rect(108, y_pos, box_w, box_h, style='FD')

        pdf.set_xy(14, y_pos + 3)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(100, 116, 139)
        pdf.cell(box_w, 5, kpis[0][0], align='C', ln=True)
        pdf.set_x(14)
        pdf.set_font("Helvetica", "B", 14)
        pdf.set_text_color(30, 58, 138)
        pdf.cell(box_w, 8, kpis[0][1], align='C')

        pdf.set_xy(108, y_pos + 3)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(100, 116, 139)
        pdf.cell(box_w, 5, kpis[1][0], align='C', ln=True)
        pdf.set_x(108)
        pdf.set_font("Helvetica", "B", 14)
        pdf.set_text_color(22, 101, 52) # Green
        pdf.cell(box_w, 8, kpis[1][1], align='C')

        # Second row of KPI boxes
        y_pos += box_h + 4
        pdf.set_fill_color(248, 250, 252) # Light gray
        pdf.set_draw_color(226, 232, 240)
        pdf.rect(14, y_pos, box_w, box_h, style='FD')
        pdf.rect(108, y_pos, box_w, box_h, style='FD')

        pdf.set_xy(14, y_pos + 3)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(100, 116, 139)
        pdf.cell(box_w, 5, kpis[2][0], align='C', ln=True)
        pdf.set_x(14)
        pdf.set_font("Helvetica", "B", 14)
        pdf.set_text_color(109, 40, 217) # Purple
        pdf.cell(box_w, 8, kpis[2][1], align='C')

        pdf.set_xy(108, y_pos + 3)
        pdf.set_font("Helvetica", "B", 9)
        pdf.set_text_color(100, 116, 139)
        pdf.cell(box_w, 5, kpis[3][0], align='C', ln=True)
        pdf.set_x(108)
        pdf.set_font("Helvetica", "B", 14)
        pdf.set_text_color(180, 83, 9) # Amber
        pdf.cell(box_w, 8, kpis[3][1], align='C')

        pdf.set_y(y_pos + box_h + 8)

        # Payment breakdown section
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(17, 24, 39)
        pdf.cell(0, 8, "Incassi per Metodo di Pagamento", ln=True)
        pdf.ln(1)

        pdf.set_fill_color(37, 99, 235)
        pdf.set_text_color(255, 255, 255)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(90, 7, "Metodo di Pagamento", border=1, fill=True)
        pdf.cell(50, 7, "Importo", border=1, fill=True, align="R")
        pdf.cell(42, 7, "% Totale", border=1, fill=True, align="R")
        pdf.ln()

        tot_rev = stats.get("total_revenue", 0.0)
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(31, 41, 55)
        for method, amt in stats.get("revenue_by_payment", {}).items():
            pct = (amt / tot_rev * 100) if tot_rev > 0 else 0.0
            pdf.cell(90, 6, f"  {method}", border=1)
            pdf.cell(50, 6, f"EUR {amt:.2f}  ", border=1, align="R")
            pdf.cell(42, 6, f"{pct:.1f}%  ", border=1, align="R")
            pdf.ln()

        pdf.ln(6)

        # Category sales breakdown section
        pdf.set_font("Helvetica", "B", 12)
        pdf.set_text_color(17, 24, 39)
        pdf.cell(0, 8, "Articoli Venduti per Categoria", ln=True)
        pdf.ln(1)

        for cat_name, items in stats.get("items_by_category", {}).items():
            if not items:
                continue
            pdf.set_fill_color(241, 245, 249)
            pdf.set_text_color(15, 23, 42)
            pdf.set_font("Helvetica", "B", 10)
            pdf.cell(140, 6, f"  {cat_name.upper()}", border=1, fill=True)
            pdf.cell(42, 6, "Quantita'  ", border=1, fill=True, align="R")
            pdf.ln()

            pdf.set_font("Helvetica", "", 9)
            pdf.set_text_color(51, 65, 85)
            for item_name, qty in items.items():
                pdf.cell(140, 5.5, f"    {item_name}", border=1)
                pdf.cell(42, 5.5, f"{qty}  ", border=1, align="R")
                pdf.ln()
            pdf.ln(2)

        os.makedirs(os.path.dirname(path), exist_ok=True)
        pdf.output(path)
